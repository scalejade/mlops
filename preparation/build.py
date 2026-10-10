#!/usr/bin/env python3
"""
Build the PBBI clause-extraction dataset: labelled clauses + source PDF -> JSONL,
in the exact request shape the app sends the extraction model.

    python preparation/build.py --tokenizer scalejade/qwen-sea-lion-v4.5-27b-pbbi

THE SERVING CONTRACT -- scalejade/perjanjian-baku-bi, branch sys-prompt @ e028187
    clause_extraction.go sends ONE call per document:

        system     prompts/extraction_system.md, byte-for-byte (go:embed, no trim)
        user       "DOCUMENT:\\n" + text
        assistant  {"clauses":[{"pasal","topic","clause_text"}, ...]}

    where text is pdftotext -layout -> keepIndonesianOnly(), ported in
    serving_text.py. Every row here is built the same way. If the app changes the
    prompt, the envelope, or the text pipeline, this file changes in the same PR,
    or the adapter is trained on an input it is never served.

ONE ROW IS A DOCUMENT -- OR A SHORTER ONE
    A labelled contract plus the 14k-token system prompt plus its answer runs 19k
    to 82k tokens; the answer alone is ~1.6x the document, because separated
    child clauses repeat their parent's text. A document that fits
    --max-row-tokens becomes one row, whole. A larger one is cut into WINDOWS:
    runs of consecutive whole sections, each sent as its own "DOCUMENT:\\n...".

    A window is still a faithful serving input -- the app sends short contracts
    in exactly this shape -- and its answer is exactly the labels of the clauses
    it contains. Three rules keep it that way:

      - a window starts at a section's header line, or at the top of the
        document, never mid-section;
      - a section whose clauses cannot be located in the PDF text ends the
        window before it, so no row carries clause text without its label;
      - sections enter a window only in document order.

DOCUMENTS
    One folder per labelled document under --source, discovered, not listed:

        <source>/<anything>/
            <one>.pdf              the contract
            <one>.xlsx             Legal's labels: No, No. Pasal, Topik Klausa, Bunyi Klausa
            <one>.gold.json        optional; the review-annotation gold version, wins over xlsx
            annotation.txt         "Split : Train | Validation | Test", plus category fields

    The split decides where a document goes, and a document is in exactly one:

        Train        train.jsonl   whole, or windowed when too long
        Validation   eval.jsonl    whole, or windowed -- the eval loss during training
        Test         test.jsonl    always whole, exactly the app's call; never trained
                                   on, never used for loss. For scoring generations.

The user side is the served PDF text, never the clean XLSX text: clean-in /
clean-out teaches the model to copy. The gap between the two IS the lesson.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import sys
import unicodedata
from dataclasses import dataclass
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("  error  openpyxl is required:  pip install openpyxl")

sys.path.insert(0, str(Path(__file__).resolve().parent))
import serving_text  # noqa: E402
from app_request import DOCUMENT_PREFIX, SYSTEM_PROMPT, user_message  # noqa: E402

HERE = Path(__file__).resolve().parent


# --------------------------------------------------------------- documents

SPLITS = {"train": "train", "validation": "eval", "test": "test"}   # split -> output file


@dataclass(frozen=True)
class Doc:
    label: str              # the folder name
    folder: Path
    pdf: Path
    labels: Path            # *.gold.json when present, else the xlsx
    split: str              # train | validation | test
    meta: dict              # annotation.txt, flattened


def read_annotation(path: Path) -> dict:
    """ "Key : Value" lines; "[Section]" lines are headings and are skipped."""
    meta = {}
    for line in path.read_text("utf-8").splitlines():
        line = line.strip()
        if not line or line.startswith("[") or ":" not in line:
            continue
        k, v = line.split(":", 1)
        meta[k.strip()] = v.strip()
    return meta


def _one(folder: Path, pattern: str, required: bool) -> Path | None:
    found = sorted(p for p in folder.glob(pattern) if not p.name.startswith("."))
    if len(found) > 1:
        raise Fail(f"{folder.name}: {len(found)} files match {pattern} — keep exactly one "
                   f"({', '.join(p.name for p in found)})")
    if not found and required:
        raise Fail(f"{folder.name}: no {pattern}")
    return found[0] if found else None


def discover(source: Path) -> list[Doc]:
    """
    Every folder under source is one document. Anything malformed fails loudly:
    a document silently skipped is a document silently missing from a split.
    """
    if not source.is_dir():
        raise Fail(f"--source is not a folder: {source}")
    docs = []
    for folder in sorted(p for p in source.iterdir()
                         if p.is_dir() and not p.name.startswith((".", "_"))):
        ann = folder / "annotation.txt"
        if not ann.exists():
            raise Fail(f"{folder.name}: no annotation.txt — every document declares its Split")
        meta = read_annotation(ann)
        split = meta.get("Split", "").lower()
        if split not in SPLITS:
            raise Fail(f"{folder.name}: Split is {meta.get('Split')!r}, "
                       f"expected one of Train, Validation, Test")
        pdf = _one(folder, "*.pdf", required=True)
        gold = _one(folder, "*.gold.json", required=False)
        labels = gold or _one(folder, "*.xlsx", required=True)
        docs.append(Doc(folder.name, folder, pdf, labels, split, meta))
    return docs


# Columns, by position: the labelled files use "No. Pasal" and "Pasal"
# interchangeably for the same column, so the header text is not trustworthy.
COL_NO, COL_PASAL, COL_TOPIC, COL_TEXT = 0, 1, 2, 3

ANCHOR_CHARS = 60      # normalised characters used to locate a clause in the PDF
MIN_ANCHOR = 24        # a clause shorter than this is too generic to anchor on
CHARS_PER_TOKEN = 3.5  # Indonesian, cl100k-ish. Only used without --tokenizer.


class Fail(Exception):
    """A problem the user needs to fix. Printed without a traceback."""


def info(msg: str) -> None:
    print(f"  {msg}")


def head(msg: str) -> None:
    print(f"\n{msg}")


# ------------------------------------------------------------------ labels

@dataclass
class Clause:
    no: str
    pasal: str
    topic: str
    text: str

    @property
    def section(self) -> str:
        return self.pasal.split("|")[0].strip()


def load_sheet(xlsx: Path, sheet: str | None = None) -> list[Clause]:
    """The named sheet, or the only one. Several sheets and no name is an error:
    picking one would be a guess about which answer key goes with the PDF."""
    book = openpyxl.load_workbook(xlsx, data_only=True)
    if sheet is None:
        if len(book.sheetnames) != 1:
            raise Fail(f"{xlsx.name} has {len(book.sheetnames)} sheets "
                       f"({', '.join(book.sheetnames)}) — name one with "
                       f"'Sheet : <name>' in annotation.txt")
        sheet = book.sheetnames[0]
    if sheet not in book.sheetnames:
        raise Fail(f"{xlsx.name} has no sheet {sheet!r} (has: {', '.join(book.sheetnames)})")
    rows = []
    for row in book[sheet].iter_rows(min_row=2, values_only=True):
        if len(row) <= COL_TEXT:
            continue
        pasal, topic, text = row[COL_PASAL], row[COL_TOPIC], row[COL_TEXT]
        if not pasal or not text:
            continue
        rows.append(Clause(no=str(row[COL_NO] or "").strip(), pasal=str(pasal).strip(),
                           topic=str(topic or "").strip(), text=str(text).strip()))
    return rows


def load_gold(path: Path) -> list[Clause]:
    data = json.loads(path.read_text("utf-8"))
    return [Clause(no=str(r.get("no", "")), pasal=str(r["pasal"]).strip(),
                   topic=str(r.get("topik", "")).strip(), text=str(r["bunyi"]).strip())
            for r in data["rows"] if r.get("pasal") and r.get("bunyi")]


def load_labels(doc: Doc) -> tuple[list[Clause], str]:
    if doc.labels.name.endswith(".gold.json"):
        clauses, src = load_gold(doc.labels), f"gold:{doc.labels.name}"
    else:
        clauses, src = load_sheet(doc.labels, doc.meta.get("Sheet")), f"xlsx:{doc.labels.name}"
    if not clauses:
        raise Fail(f"{doc.label}: {doc.labels.name} has no labelled rows")
    return clauses, src


# -------------------------------------------------------------- alignment

_WS = re.compile(r"\s+")


def _fold(ch: str) -> str:
    """Curly quotes, dashes and NBSP differ between the labels and the PDF dump."""
    return {
        "‘": "'", "’": "'", "“": '"', "”": '"',
        "–": "-", "—": "-", " ": " ",
    }.get(ch, ch)


def normalise(s: str) -> str:
    s = unicodedata.normalize("NFKC", s)
    s = "".join(_fold(c) for c in s)
    return _WS.sub(" ", s).strip().lower()


class Aligned:
    """
    A normalised view of the document text that maps any position in the
    normalised string back to an offset in the original, so a match found on
    whitespace-insensitive text can be sliced out of the text the model will see.
    """

    def __init__(self, raw: str):
        self.raw = raw
        chars: list[str] = []
        index: list[int] = []
        prev_space = True  # leading whitespace is dropped, as in normalise()
        for i, ch in enumerate(unicodedata.normalize("NFKC", raw)):
            ch = _fold(ch)
            if ch.isspace():
                if not prev_space:
                    chars.append(" ")
                    index.append(i)
                prev_space = True
                continue
            chars.append(ch.lower())
            index.append(i)
            prev_space = False
        while chars and chars[-1] == " ":
            chars.pop()
            index.pop()
        self.norm = "".join(chars)
        self.index = index

    def find(self, needle: str, start_at: int = 0) -> int:
        return self.norm.find(needle, start_at)

    def orig(self, norm_pos: int) -> int:
        return self.index[min(norm_pos, len(self.index) - 1)]


def anchors(text: str) -> tuple[str, str] | None:
    """Head and tail slices of a clause, long enough to be unique in a contract."""
    n = normalise(text)
    if len(n) < MIN_ANCHOR:
        return None
    return n[:ANCHOR_CHARS], n[-ANCHOR_CHARS:]


def line_start(raw: str, pos: int) -> int:
    nl = raw.rfind("\n", 0, pos)
    return 0 if nl < 0 else nl + 1


def line_end(raw: str, pos: int) -> int:
    nl = raw.find("\n", pos)
    return len(raw) if nl < 0 else nl


@dataclass
class Section:
    header: str
    clauses: list[Clause]
    start: int = -1        # first clause's line, in the served text
    end: int = -1          # end of the last clause's line
    error: str = ""
    # Set on the 2nd, 3rd, ... piece of a section split by split_oversized():
    # the piece starts mid-section, so a window opening on it is given the
    # section's own header line(s) from the PDF first.
    parent: "Section | None" = None

    @property
    def ok(self) -> bool:
        return not self.error


def group_sections(clauses: list[Clause]) -> list[Section]:
    """Consecutive clauses sharing a first pasal segment. Consecutive, not
    grouped by key: a header that recurs later is a separate span."""
    out: list[Section] = []
    for c in clauses:
        if out and out[-1].header == c.section:
            out[-1].clauses.append(c)
        else:
            out.append(Section(c.section, [c]))
    return out


def locate(sec: Section, doc: Aligned, cursor: int, min_coverage: float) -> int:
    """
    Find one section's span, searching forward from cursor -- contracts repeat
    boilerplate, and a backwards match would splice unrelated parts together.
    Sets sec.start/end, or sec.error. Returns the new cursor.

    A section whose boundaries are a guess is worse than a section that is
    dropped: the model would be trained to emit clauses that are not in its
    input, or to skip ones that are.
    """
    first, last = sec.clauses[0], sec.clauses[-1]
    a, b = anchors(first.text), anchors(last.text)
    if a is None:
        sec.error = "first clause too short to anchor"
        return cursor
    if b is None:
        sec.error = "last clause too short to anchor"
        return cursor
    head_pos = doc.find(a[0], cursor)
    if head_pos < 0:
        sec.error = "first clause not found after the previous section"
        return cursor
    tail_pos = doc.find(b[1], head_pos)
    if tail_pos < 0:
        sec.error = "last clause not found after the first"
        return cursor

    # How many of the clauses in between are actually there, in order. A section
    # that scores low is misaligned even though its two ends matched.
    found, scan, anchored = 0, head_pos, 0
    for c in sec.clauses:
        ac = anchors(c.text)
        if ac is None:
            continue
        anchored += 1
        p = doc.find(ac[0], scan)
        if 0 <= p <= tail_pos + ANCHOR_CHARS:
            found += 1
            scan = p
    coverage = found / max(1, anchored)
    if coverage < min_coverage:
        sec.error = f"coverage {coverage:.0%} < {min_coverage:.0%}"
        return cursor

    sec.start = line_start(doc.raw, doc.orig(head_pos))
    sec.end = line_end(doc.raw, doc.orig(tail_pos + len(b[1]) - 1))
    return tail_pos


def split_oversized(sec: Section, doc: Aligned, min_coverage: float) -> list[Section]:
    """
    Break a located section that is too long for one row on its second pasal
    segment ("2. Layanan|2.2|..." -> the 2.2 piece). Deterministic: no guessing
    where a boundary goes, only the numbering the labellers recorded. Each piece
    is located inside the section's own span, in order.
    """
    pieces: list[Section] = []
    for c in sec.clauses:
        parts = c.pasal.split("|")
        key = parts[1].strip() if len(parts) > 1 else ""
        name = f"{sec.header}|{key}" if key else sec.header
        if pieces and pieces[-1].header == name:
            pieces[-1].clauses.append(c)
        else:
            pieces.append(Section(name, [c]))
    if len(pieces) <= 1:
        return [sec]

    cursor = next(i for i, o in enumerate(doc.index) if o >= sec.start)
    for k, p in enumerate(pieces):
        cursor = locate(p, doc, cursor, min_coverage)
        if p.ok and (p.start < sec.start or p.end > sec.end or
                     (k and pieces[k - 1].ok and p.start < pieces[k - 1].end)):
            p.error = "piece located outside or out of order within its section"
        p.parent = sec
    return pieces


def leaks(body: str, inside: list[Clause], every: list[Clause]) -> list[Clause]:
    """
    Labelled clauses whose text is in the window but not in its answer. Any
    is a reason to drop the window: the row would teach the model to skip a
    clause. Catches multi-column PDFs, where pdftotext -layout puts two
    unrelated clauses on one line, and any other misalignment.
    """
    n = normalise(body)
    mine = {id(c) for c in inside}
    # Text the window's own labels account for. A phrase another clause shares
    # with one of these -- boilerplate, inherited parent text -- is covered.
    own = "\n".join(normalise(c.text) for c in inside)
    out = []
    for c in every:
        if id(c) in mine:
            continue
        a = anchors(c.text)
        if a and any(p in n and p not in own for p in a):
            out.append(c)
    return out


def header_line(sec: Section, doc: Aligned, lo: int) -> int:
    """
    Where a window starting at this section begins: the section's header line,
    searched between the previous located text (lo) and the first clause. -1 if
    the header text is not there -- e.g. "1" for an un-headed opening clause, or
    a label whose header is spelled differently from the PDF.
    """
    h = normalise((sec.parent or sec).header)
    if len(h) < 3:
        return -1
    lo_n = next((i for i, o in enumerate(doc.index) if o >= lo), len(doc.index))
    hi_n = next((i for i, o in enumerate(doc.index) if o >= sec.start), len(doc.index))
    p = doc.norm.rfind(h[:ANCHOR_CHARS], lo_n, hi_n)
    return -1 if p < 0 else line_start(doc.raw, doc.orig(p))


# ---------------------------------------------------------------- assembly

def answer(clauses: list[Clause]) -> str:
    """
    The assistant turn. Compact, this exact key order, nothing around it --
    extraction_system.md section 11 asks for the bare object, and the model
    reproduces formatting along with content. ensure_ascii=False keeps
    Indonesian readable and the token count honest.
    """
    return json.dumps(
        {"clauses": [{"pasal": c.pasal, "topic": c.topic, "clause_text": c.text}
                     for c in clauses]},
        ensure_ascii=False, separators=(",", ":"),
    )


def make_row(system: str, text: str, clauses: list[Clause]) -> dict:
    return {"messages": [
        {"role": "system", "content": system},
        {"role": "user", "content": user_message(text)},
        {"role": "assistant", "content": answer(clauses)},
    ]}


class Counter:
    def __init__(self, tokenizer: str | None, system: str):
        self.tok = None
        if tokenizer:
            # `tokenizers` alone (no transformers, no torch) is enough for counts,
            # so a laptop needs only preparation/requirements.txt.
            from tokenizers import Tokenizer
            if Path(tokenizer).is_file():
                t = Tokenizer.from_file(tokenizer)
            else:
                repo, _, rev = tokenizer.partition("@")
                t = Tokenizer.from_pretrained(repo, revision=rev or "main",
                                              token=os.environ.get("HF_TOKEN"))
            self.tok = lambda s: len(t.encode(s, add_special_tokens=False).ids)
        self.system = self(system)

    def __call__(self, s: str) -> int:
        return self.tok(s) if self.tok else int(len(s) / CHARS_PER_TOKEN)

    def row(self, r: dict) -> int:
        m = r["messages"]
        return self.system + self(m[1]["content"]) + self(m[2]["content"])


def build_doc(doc: Doc, system: str, count: Counter,
              max_row: int, window: int, min_coverage: float,
              rows: list[dict], dropped: list[dict], whole_only: bool = False) -> None:
    """Append doc's rows. whole_only (the Test split): one row, the whole
    document, whatever its length -- exactly the call the app makes."""
    clauses, labels_from = load_labels(doc)
    text = serving_text.served_document_text(doc.pdf)
    if serving_text.is_english_document(text):
        dropped.append({"doc": doc.label, "section": "*", "clauses": len(clauses),
                        "reason": "the app rejects it as an English document"})
        info(f"{doc.label:<30} skipped — the app would reject it as English")
        return

    # Whole document, when it fits: the exact call the app makes.
    whole = make_row(system, text, clauses)
    n = count.row(whole)
    if n <= max_row or whole_only:
        whole["_meta"] = {"doc": doc.label, "kind": "document", "labels": labels_from,
                          "sections": "*", "clauses": len(clauses), "tokens": n,
                          "annotation": doc.meta}
        rows.append(whole)
        info(f"{doc.label:<30} {len(clauses):>4} klausa  {labels_from:<36} "
             f"-> 1 document row ({n:,} tok)")
        return

    # Otherwise windows of consecutive, located sections.
    aligned = Aligned(text)
    units: list[Section] = []
    cursor = 0
    for s in group_sections(clauses):
        cursor = locate(s, aligned, cursor, min_coverage)
        if s.ok and count(answer(s.clauses)) + count(text[s.start:s.end]) + count.system > window:
            units.extend(split_oversized(s, aligned, min_coverage))
        else:
            units.append(s)
    for u in units:
        if not u.ok:
            dropped.append({"doc": doc.label, "section": u.header,
                            "clauses": len(u.clauses), "reason": u.error})

    made = 0
    win: list[Section] = []
    win_start = 0
    lead = ""             # header line(s) of the section a mid-section window opens in
    head_of: dict[int, int] = {}   # id(section) -> where its header region begins

    def flush(end_of_doc: bool) -> None:
        nonlocal made
        if not win:
            return
        end = len(text) if end_of_doc else win[-1].end
        body = (lead + text[win_start:end]).strip()
        cl = [c for s in win for c in s.clauses]
        name = " + ".join(s.header for s in win)
        leaked = leaks(body, cl, clauses)
        if leaked:
            dropped.append({"doc": doc.label, "section": name, "clauses": len(cl),
                            "reason": f"window text contains {len(leaked)} labelled clause(s) "
                                      f"outside it (e.g. {leaked[0].pasal}) — multi-column "
                                      f"layout or misalignment"})
            return
        r = make_row(system, body, cl)
        n = count.row(r)
        if n > max_row:
            dropped.append({"doc": doc.label, "section": name,
                            "clauses": len(cl), "reason": f"{n:,} tok > --max-row-tokens"})
            return
        r["_meta"] = {"doc": doc.label, "kind": "window", "labels": labels_from,
                      "sections": [s.header for s in win], "clauses": len(cl), "tokens": n,
                      "annotation": doc.meta}
        rows.append(r)
        made += 1

    def continues(u: Section) -> bool:
        """A piece of a split section other than its first."""
        return u.parent is not None and u.parent.clauses[0] is not u.clauses[0]

    def opening(i: int, u: Section, lo: int) -> tuple[int, str]:
        """Where a window that opens on unit i begins, and what it is prefixed
        with. lo is the end of the last located unit before it. (-1, "") when
        there is no honest place to begin it."""
        if continues(u):
            # Mid-section piece: the section's header region, then the piece.
            h = head_of.get(id(u.parent), -1)
            if h < 0:
                return -1, ""
            return u.start, text[h:u.parent.start].strip("\n") + "\n"
        if i == 0:
            return 0, ""                    # the top of the document
        if units[i - 1].ok:
            return units[i - 1].end, ""     # the gap is this section's header
        return header_line(u, aligned, lo), ""

    last_end = 0          # end of the last located unit
    for i, u in enumerate(units):
        sec = u.parent or u
        if id(sec) not in head_of:
            # First unit seen of a section: remember where its header region
            # starts, for any later piece of it that has to open a window.
            head_of[id(sec)] = (0 if i == 0 else units[i - 1].end if units[i - 1].ok
                                else header_line(sec, aligned, last_end))

        if not u.ok:
            flush(False)
            win = []
            continue
        lo, last_end = last_end, u.end

        if win:
            trial = make_row(system, (lead + text[win_start:u.end]).strip(),
                             [c for x in win + [u] for c in x.clauses])
            if count.row(trial) <= window:
                win.append(u)
                continue
            flush(False)
            win = []

        start, lead_text = opening(i, u, lo)
        if start < 0:
            dropped.append({"doc": doc.label, "section": u.header, "clauses": len(u.clauses),
                            "reason": "follows an unlocated section and its header line "
                                      "was not found to start a window on"})
            continue
        win, win_start, lead = [u], start, lead_text

    flush(bool(win) and win[-1] is units[-1])
    info(f"{doc.label:<30} {len(clauses):>4} klausa  {labels_from:<36} "
         f"-> {made} window rows ({n:,} tok whole)")


# -------------------------------------------------------------------- main

def write_jsonl(path: Path, rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8") as fh:
        for r in rows:
            fh.write(json.dumps({"messages": r["messages"]}, ensure_ascii=False) + "\n")


def write_sample(path: Path, rows: list[dict], n: int = 8) -> None:
    """
    The only file here that git tracks, so nothing quotable from a client contract
    goes in it: every content field is cut to 120 characters. It exists to show a
    reviewer the shape of a row, not to be trained on.
    """
    step = max(1, len(rows) // n)
    picked = rows[::step][:n]
    with path.open("w", encoding="utf-8") as fh:
        for r in picked:
            msgs = [{"role": m["role"], "content": m["content"][:120] + " […TRUNCATED]"}
                    for m in r["messages"]]
            fh.write(json.dumps({"messages": msgs}, ensure_ascii=False) + "\n")


def summarise(name: str, rows: list[dict], dropped: list[dict]) -> None:
    head(f"{name} summary")
    if not rows:
        info("no rows")
        return
    lengths = sorted(r["_meta"]["tokens"] for r in rows)
    kinds = {k: sum(1 for r in rows if r["_meta"]["kind"] == k) for k in ("document", "window")}
    info(f"rows            {len(rows)}   ({kinds['document']} whole documents, "
         f"{kinds['window']} windows)")
    info(f"clauses         {sum(r['_meta']['clauses'] for r in rows)}")
    info(f"tokens/row      median {lengths[len(lengths)//2]:,}   max {lengths[-1]:,}")
    info(f"dropped         {len(dropped)} section(s), "
         f"{sum(d['clauses'] for d in dropped)} clauses")
    for d in dropped:
        sec = d["section"] if isinstance(d["section"], str) else str(d["section"])
        info(f"   - {d['doc']:<28} {sec[:38]:<40} {d['clauses']:>3} kl  {d['reason']}")


def composition(docs: list[Doc]) -> None:
    """Documents per split, by the annotation fields that matter for coverage.
    A split that is all one Kombinasi or Produk measures only that slice."""
    head("composition")
    for field in ("Kombinasi", "Produk"):
        values = sorted({d.meta.get(field, "?") for d in docs})
        width = max(len(v) for v in values + [field])
        info(f"{field:<{width}}  " + "  ".join(f"{n:>5}" for n in SPLITS.values()))
        for v in values:
            cells = [sum(1 for d in docs if d.split == sp and d.meta.get(field, "?") == v)
                     for sp in SPLITS]
            info(f"{v:<{width}}  " + "  ".join(f"{c:>5}" for c in cells))
        info("")


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--source", type=Path, default=HERE / "raw",
                    help="one folder per document: a PDF, its labels (xlsx or "
                         "*.gold.json) and annotation.txt (default: preparation/raw)")
    ap.add_argument("--out", type=Path, default=HERE / "data",
                    help="where train.jsonl / eval.jsonl / test.jsonl / sample.jsonl / "
                         "build-report.json go (default: preparation/data)")
    ap.add_argument("--max-row-tokens", type=int, default=30000,
                    help="a Train/Validation document longer than this is windowed; no "
                         "such row is longer. Keep it under model.max_seq_length with room "
                         "for the chat template (default 30000, for max_seq_length 32768). "
                         "Test rows are always whole.")
    ap.add_argument("--window-tokens", type=int, default=24000,
                    help="target length of a window row, system prompt included "
                         "(default 24000 — the prompt alone is ~14k)")
    ap.add_argument("--min-coverage", type=float, default=0.8,
                    help="drop a section when fewer than this fraction of its clauses "
                         "are found, in order, inside the located span (default 0.8)")
    ap.add_argument("--tokenizer", help="HF repo id (repo or repo@revision; private repos "
                                        "read HF_TOKEN), or a local tokenizer.json, for exact "
                                        "token counts; without it lengths are estimated "
                                        f"at {CHARS_PER_TOKEN} chars/token")
    ap.add_argument("--dry-run", action="store_true", help="report only, write nothing")
    args = ap.parse_args()

    # read_text, not .strip(): go:embed sends the file byte-for-byte.
    system = SYSTEM_PROMPT.read_text("utf-8")
    count = Counter(args.tokenizer, system)
    info(f"tokenizer       {args.tokenizer or f'none — ESTIMATED at {CHARS_PER_TOKEN} chars/token'}")
    info(f"system prompt   {SYSTEM_PROMPT.name}, {len(system):,} chars, {count.system:,} tok")
    info(f"text pipeline   {serving_text.SOURCE}")

    docs = discover(args.source)
    head("documents")
    for d in docs:
        info(f"{d.label:<30} {d.split:<11} {d.labels.name}")
    if not docs:
        raise Fail(f"no document folders under {args.source}")
    composition(docs)

    # A document in two splits makes the eval or test number a measure of recall.
    # Folder names can differ for the same contract, so compare the PDFs' bytes.
    seen: dict[str, Doc] = {}
    for d in docs:
        h = hashlib.sha256(d.pdf.read_bytes()).hexdigest()
        if h in seen and seen[h].split != d.split:
            raise Fail(f"the same PDF is in {seen[h].split} ({seen[h].label}) and "
                       f"{d.split} ({d.label})")
        seen[h] = d

    out: dict[str, tuple[list[dict], list[dict]]] = {}
    for split, name in SPLITS.items():
        head(f"{name}  ({split})")
        rows: list[dict] = []
        dropped: list[dict] = []
        members = [d for d in docs if d.split == split]
        if not members:
            info("no documents")
        for d in members:
            build_doc(d, system, count, args.max_row_tokens, args.window_tokens,
                      args.min_coverage, rows, dropped, whole_only=(split == "test"))
        out[name] = (rows, dropped)

    for name in SPLITS.values():
        summarise(name, *out[name])

    if args.dry_run:
        head("dry run — nothing written")
        return

    head("written")
    args.out.mkdir(parents=True, exist_ok=True)
    for name, (rows, _) in out.items():
        path = args.out / f"{name}.jsonl"
        if rows:
            write_jsonl(path, rows)
            info(f"{path}   {len(rows)} rows")
        elif path.exists():
            # A file left from an earlier build would be trained on as if current.
            path.replace(path.with_suffix(".jsonl.prev"))
            info(f"{path}   NOT written — no rows; the previous build's file moved to "
                 f"{path.name}.prev")
        else:
            info(f"{path}   NOT written — no rows")
    train = out["train"][0]
    if train:
        write_sample(args.out / "sample.jsonl", train)
        info(f"{args.out / 'sample.jsonl'}   8 rows, truncated, safe to commit")
    (args.out / "build-report.json").write_text(json.dumps({
        "serving_contract": {"system_prompt": SYSTEM_PROMPT.name,
                             "user_prefix": DOCUMENT_PREFIX,
                             "text_pipeline": serving_text.SOURCE},
        "documents": [{"folder": d.label, "split": d.split, "pdf": d.pdf.name,
                       "labels": d.labels.name, "annotation": d.meta} for d in docs],
        **{name: {"rows": len(rows), "dropped": dropped, "meta": [r["_meta"] for r in rows]}
           for name, (rows, dropped) in out.items()},
    }, ensure_ascii=False, indent=2), encoding="utf-8")
    info(f"{args.out / 'build-report.json'}")


if __name__ == "__main__":
    try:
        main()
    except (Fail, serving_text.ExtractError) as e:
        sys.exit(f"\n  error  {e}\n")
