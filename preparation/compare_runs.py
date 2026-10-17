#!/usr/bin/env python3
"""
Put two evaluation runs side by side, clause by clause, in one workbook: Legal's
labels next to what each model generated for the same contract.

    python preparation/compare_runs.py \\
        preparation/data/scores/test-v0.1.3-t0 "v0.1.3 (fine-tuned)" \\
        preparation/data/scores/test-gpt-4.1-mini-t0 "gpt-4.1-mini" \\
        --out docs/pbbi/reports/pbbi-extraction-clauses-v0.1.3-vs-gpt-4.1-mini.xlsx

Each run directory is evaluate.py's --out: report.json plus replies/. Both runs
must be over the same data file (test.jsonl by default). Matching is score.py's:
both sides through the app's parse_reply + finalize, then one-to-one at >= 0.8
word overlap, closest pairs first — the numbers equal evaluate.py's.

THE WORKBOOK QUOTES THE CONTRACTS. Keep it out of git (docs/pbbi/reports/*.xlsx
is gitignored) and share it the way the source contracts are shared.
"""

from __future__ import annotations

import argparse
import json
import sys
from pathlib import Path

import openpyxl
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import app_request as app  # noqa: E402
import score  # noqa: E402

NEAREST_SHOWN = 0.3     # an unmatched label shows its closest unmatched clause above this
FILL = {k: PatternFill("solid", fgColor=c) for k, c in {
    "match": "D9EAD3", "miss": "F4CCCC", "extra": "FCE5CD", "head": "0B3D5C",
    "m1": "DCE6F1", "m2": "EDE4F3", "legal": "F3F3F3"}.items()}
BOLD, WHITE = Font(bold=True), Font(bold=True, color="FFFFFF")
WRAP = Alignment(wrap_text=True, vertical="top")


def clean(s) -> str:
    return ILLEGAL_CHARACTERS_RE.sub("", str(s)) if s is not None else ""


def jaccard(a: frozenset, b: frozenset) -> float:
    u = len(a | b)
    return len(a & b) / u if u else 1.0


def pair(gold: list, pred: list):
    """score.match_clauses, keeping the pairs: gold index -> (pred index, similarity),
    plus the closest unmatched prediction for each unmatched label."""
    gw = [score.clause_words(c) for c in gold]
    pw = [score.clause_words(c) for c in pred]
    cand = sorted(((jaccard(g, p), i, j) for i, g in enumerate(gw) for j, p in enumerate(pw)
                   if jaccard(g, p) >= score.CLAUSE_MATCH), key=lambda x: (-x[0], x[1], x[2]))
    g2p, used = {}, set()
    for s, i, j in cand:
        if i not in g2p and j not in used:
            g2p[i] = (j, s)
            used.add(j)
    near = {}
    for i, g in enumerate(gw):
        if i not in g2p:
            best = max(((jaccard(g, pw[j]), j) for j in range(len(pred)) if j not in used),
                       default=(0.0, None))
            if best[1] is not None and best[0] >= NEAREST_SHOWN:
                near[i] = (best[1], best[0])
    return g2p, near, used


def load_run(path: Path, label: str) -> dict:
    report = json.loads((path / "report.json").read_text("utf-8"))
    models = list(report["summary"])
    if len(models) != 1:
        sys.exit(f"{path}: has {len(models)} models ({', '.join(models)}); pass a single-model run")
    model = models[0]
    results = {r["row"]: r for r in report["results"]}
    replies = {}
    for i in results:
        f = path / "replies" / f"{model.replace('/', '__')}__row{i}.txt"
        replies[i] = f.read_text("utf-8") if f.exists() else ""
    return {"label": label, "model": model, "report": report, "results": results,
            "replies": replies, "temperature": report.get("temperature")}


def parsed(run: dict, i: int) -> list:
    r = run["results"][i]
    if "error" in r or r.get("finish") == "length":
        return []
    try:
        return app.finalize(app.parse_reply(run["replies"][i]))[0]
    except app.AppError:
        return []


def header(ws, row: int, values: list, fills: list | None = None) -> None:
    for k, v in enumerate(values, 1):
        c = ws.cell(row, k, v)
        c.font, c.alignment = WHITE, WRAP
        c.fill = (fills[k - 1] if fills else None) or FILL["head"]


def widths(ws, w: list) -> None:
    for k, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(k)].width = x


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("run1", type=Path)
    ap.add_argument("label1")
    ap.add_argument("run2", type=Path)
    ap.add_argument("label2")
    ap.add_argument("--data", type=Path, default=HERE / "data" / "test.jsonl")
    ap.add_argument("--out", type=Path, required=True)
    args = ap.parse_args()

    rows = [json.loads(l) for l in args.data.read_text("utf-8").splitlines() if l.strip()]
    try:
        meta = json.loads((args.data.parent / "build-report.json").read_text("utf-8"))[args.data.stem]["meta"]
    except (OSError, KeyError):
        meta = [{"doc": f"row {i}", "annotation": {}} for i in range(len(rows))]
    runs = [load_run(args.run1, args.label1), load_run(args.run2, args.label2)]
    picked = sorted(set(runs[0]["results"]) & set(runs[1]["results"]))

    # ---- match every contract once, for every sheet --------------------------
    per = {}
    for i in picked:
        gold = app.finalize(app.parse_reply(rows[i]["messages"][2]["content"]))[0]
        sides = []
        for run in runs:
            pred = parsed(run, i)
            g2p, near, used = pair(gold, pred)
            r = run["results"][i]
            sides.append({"pred": pred, "g2p": g2p, "near": near, "used": used, "result": r,
                          "matched": len(g2p), "returned": len(pred)})
        per[i] = {"gold": gold, "sides": sides}

    def prf(m, lab, ret):
        R = m / lab if lab else 0.0
        P = m / ret if ret else 0.0
        return R, P, (2 * P * R / (P + R) if P + R else 0.0)

    wb = openpyxl.Workbook()

    # ---- Ringkasan -------------------------------------------------------------
    ws = wb.active
    ws.title = "Ringkasan"
    ws.append(["PBBI clause extraction — clause-by-clause comparison"])
    ws["A1"].font = Font(bold=True, size=14, color="0B3D5C")
    ws.append([f"Data: {args.data.name}, {len(picked)} contracts never used in training. "
               f"Both runs used the app's exact request and the app's own parser."])
    ws.append([])
    header(ws, 4, ["", runs[0]["label"], runs[1]["label"]])
    lab = sum(len(per[i]["gold"]) for i in picked)
    lines = []
    for k, run in enumerate(runs):
        m = sum(per[i]["sides"][k]["matched"] for i in picked)
        ret = sum(per[i]["sides"][k]["returned"] for i in picked)
        R, P, F = prf(m, lab, ret)
        res = [run["results"][i] for i in picked]
        lines.append({
            "Model": run["model"],
            "Temperature": "0 (deterministic)" if run["temperature"] == 0 else
                           ("default (1.0, sampled)" if run["temperature"] is None else run["temperature"]),
            "Legal's clauses (labelled)": lab,
            "Clauses returned": ret,
            "Matched (≥ 80% same words)": m,
            "Recall — Legal's clauses found": round(R, 3),
            "Precision — returned clauses that are right": round(P, 3),
            "F1": round(F, 3),
            "Returned ÷ labelled": round(ret / lab, 2) if lab else 0,
            "Valid JSON": f"{sum(1 for r in res if r.get('json_valid'))}/{len(res)}",
            "Finished on its own": f"{sum(1 for r in res if r.get('finished'))}/{len(res)}",
            "Output tokens": sum(r.get("output_tokens", 0) for r in res),
            "Longest contract (min)": round(max(r.get("seconds", 0) for r in res) / 60, 1),
        })
    for key in lines[0]:
        ws.append([key, lines[0][key], lines[1][key]])
        ws.cell(ws.max_row, 1).font = BOLD
    ws.append([])
    ws.append(["How to read the sheets"])
    ws.cell(ws.max_row, 1).font = Font(bold=True, color="0B3D5C")
    for a, b in [
        ("Per kontrak", "Scores per contract for both models, and which one has the higher F1."),
        ("Perbandingan klausa", "One row per clause Legal labelled, with each model's clause beside it. "
                                "After each contract's labelled clauses come the clauses a model returned "
                                "that match none of Legal's."),
        (f"{runs[0]['label']} / {runs[1]['label']} (mentah)",
         "Each model's output exactly as it was generated, in order, after the app's parser."),
        ("Cocok (green)", "The model's clause shares ≥ 80% of its words with Legal's clause (one-to-one)."),
        ("Tidak cocok (red)", "No matching clause. When the model has a clause with ≥ 30% of the same "
                              "words, it is shown as the closest one, with its similarity."),
        ("Hanya di model (orange)", "A clause the model returned that matches none of Legal's: "
                                    "split differently, merged, or not labelled by Legal."),
        ("Kemiripan", "Word overlap with Legal's clause, 0–1 (Jaccard on the app's normalised words)."),
    ]:
        ws.append([a, b])
        ws.cell(ws.max_row, 1).font = BOLD
        ws.cell(ws.max_row, 2).alignment = WRAP
    widths(ws, [44, 70, 70])

    # ---- Per kontrak -----------------------------------------------------------
    ws = wb.create_sheet("Per kontrak")
    cols = ["Row", "Contract", "Penyelenggara", "Produk", "Labelled"]
    for run in runs:
        cols += [f"{run['label']} — returned", "matched", "recall", "precision", "F1", "minutes"]
    cols += ["Higher F1"]
    header(ws, 1, cols)
    for i in sorted(picked, key=lambda i: -prf(per[i]["sides"][0]["matched"], len(per[i]["gold"]),
                                                per[i]["sides"][0]["returned"])[2]):
        an = meta[i].get("annotation", {})
        line = [i, meta[i]["doc"], an.get("Kombinasi", ""), an.get("Produk", ""), len(per[i]["gold"])]
        f1s = []
        for side in per[i]["sides"]:
            R, P, F = prf(side["matched"], len(per[i]["gold"]), side["returned"])
            f1s.append(F)
            line += [side["returned"], side["matched"], round(R, 3), round(P, 3), round(F, 3),
                     round(side["result"].get("seconds", 0) / 60, 1)]
        line.append(runs[0]["label"] if f1s[0] > f1s[1] else runs[1]["label"] if f1s[1] > f1s[0] else "tie")
        ws.append(line)
    widths(ws, [5, 40, 24, 30, 9] + [12, 9, 8, 9, 8, 8] * 2 + [22])
    ws.freeze_panes = "C2"

    # ---- Perbandingan klausa ---------------------------------------------------
    ws = wb.create_sheet("Perbandingan klausa")
    cols = ["Contract", "No", "No. Pasal (Legal)", "Topik Klausa (Legal)", "Bunyi Klausa (Legal)"]
    fills = [FILL["head"]] * 2 + [None] * 3
    for k, run in enumerate(runs):
        cols += [f"{run['label']}: status", "Kemiripan", "No. Pasal", "Topik Klausa", "Bunyi Klausa"]
        fills += [None] * 5
    header(ws, 1, cols)
    for k in range(2):
        for c in range(6 + 5 * k, 11 + 5 * k):
            ws.cell(1, c).fill = PatternFill("solid", fgColor=("2E5C8A" if k == 0 else "5E3C7A"))

    def side_cells(side, gi):
        if gi in side["g2p"]:
            j, s = side["g2p"][gi]
            c = side["pred"][j]
            return ["Cocok", round(s, 2), c.pasal, c.topic, c.clause_text], "match"
        if gi in side["near"]:
            j, s = side["near"][gi]
            c = side["pred"][j]
            return ["Tidak cocok (terdekat)", round(s, 2), c.pasal, c.topic, c.clause_text], "miss"
        return ["Tidak ditemukan", "", "", "", ""], "miss"

    for i in picked:
        gold, sides = per[i]["gold"], per[i]["sides"]
        doc = meta[i]["doc"]
        for gi, g in enumerate(gold):
            line = [doc, gi + 1, g.pasal, g.topic, g.clause_text]
            kinds = []
            for side in sides:
                cells, kind = side_cells(side, gi)
                line += cells
                kinds.append(kind)
            ws.append([clean(x) for x in line])
            r = ws.max_row
            for c in range(1, 6):
                ws.cell(r, c).fill = FILL["legal"]
            for k, kind in enumerate(kinds):
                for c in range(6 + 5 * k, 11 + 5 * k):
                    ws.cell(r, c).fill = FILL[kind]
        for k, side in enumerate(sides):
            for j, c in enumerate(side["pred"]):
                if j in side["used"]:
                    continue
                best = max((jaccard(score.clause_words(c), score.clause_words(g)) for g in gold), default=0)
                line = [doc, "", "", "", ""]
                blank = ["", "", "", "", ""]
                mine = [f"Hanya di {runs[k]['label']}", round(best, 2), c.pasal, c.topic, c.clause_text]
                line += (mine + blank) if k == 0 else (blank + mine)
                ws.append([clean(x) for x in line])
                r = ws.max_row
                for col in range(6 + 5 * k, 11 + 5 * k):
                    ws.cell(r, col).fill = FILL["extra"]
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = WRAP
    widths(ws, [26, 5, 26, 26, 60] + [16, 9, 24, 24, 60] * 2)
    ws.freeze_panes = "F2"
    ws.auto_filter.ref = ws.dimensions

    # ---- raw output per model --------------------------------------------------
    for k, run in enumerate(runs):
        ws = wb.create_sheet(f"{run['label'][:22]} (mentah)")
        header(ws, 1, ["Contract", "#", "No. Pasal", "Topik Klausa", "Bunyi Klausa", "Matches Legal?"])
        for i in picked:
            side = per[i]["sides"][k]
            for j, c in enumerate(side["pred"]):
                ws.append([clean(x) for x in (meta[i]["doc"], j + 1, c.pasal, c.topic, c.clause_text,
                                                "Ya" if j in side["used"] else "Tidak")])
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = WRAP
        widths(ws, [26, 5, 30, 30, 80, 10])
        ws.freeze_panes = "C2"
        ws.auto_filter.ref = ws.dimensions

    args.out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(args.out)
    print(f"  written    {args.out}")
    for k, run in enumerate(runs):
        print(f"  {run['label']:<24} F1 {lines[k]['F1']}  recall {lines[k]['Recall — Legal' + chr(39) + 's clauses found']}"
              f"  precision {lines[k]['Precision — returned clauses that are right']}")


if __name__ == "__main__":
    main()
